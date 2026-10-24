#!/usr/bin/env python3
"""제안서 PPTX의 슬라이드별 구조와 시각요소를 수치로 뽑는다.

글자 크기는 실제로 보이는 크기를 추정한다: run에 직접 지정된 값이 없으면
도형 → 레이아웃 → 마스터 → 프레젠테이션 기본값 순으로 상속을 따라가고,
'도형에 맞춰 텍스트 크기 조정'(normAutofit fontScale)으로 줄어든 비율도 반영한다.

usage:
    python analyze_pptx.py deck.pptx [-o analysis.json] [--md analysis.md]
        [--min-body 12] [--max-chars 400]
"""
import argparse
import json
import re
import statistics
import sys
from collections import Counter

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
except ImportError:
    sys.exit("python-pptx가 필요합니다: pip install python-pptx")

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}
DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"

TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
FOOTER_TYPES = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.SLIDE_NUMBER}

SOURCE_RE = re.compile(r"(출처|자료\s*[:：]|참고\s*[:：]|인용|Source|Sources|Ref\.|※)", re.I)
# 출처가 필요한 통계성 수치만 잡는다. 연도·기간·인원 같은 일정/조직 숫자는 제외.
NUMBER_RE = re.compile(r"\d[\d,]*(?:\.\d+)?\s*(?:%|퍼센트|배|조|억|만|천억|원|달러|\$|USD|KRW|개사|개국)")
PLACEHOLDER_TEXT_RE = re.compile(
    r"(TBD|TBA|추후\s*작성|작성\s*예정|내용\s*입력|텍스트를\s*입력|Click to add|Lorem ipsum|xxx|○○○)", re.I
)


# ---------------------------------------------------------------- font size --

def _lvl_sz(lst_el, lvl):
    """lstStyle/titleStyle/bodyStyle 같은 목록 스타일에서 lvl의 기본 크기(pt)."""
    if lst_el is None:
        return None
    el = lst_el.find(f"a:lvl{lvl + 1}pPr/a:defRPr", NS)
    if el is not None and el.get("sz"):
        return int(el.get("sz")) / 100
    return None


def _txbody_lst(shape):
    tx = getattr(shape, "_element", None)
    if tx is None:
        return None
    body = tx.find(".//p:txBody", NS)
    if body is None:
        body = tx.find(".//a:txBody", NS)
    return None if body is None else body.find("a:lstStyle", NS)


def _ph_type(shape):
    try:
        return shape.placeholder_format.type
    except (ValueError, AttributeError):
        return None


def _inherited_chain(shape, slide, prs):
    """run에 크기가 없을 때 차례로 찾아볼 목록 스타일들."""
    chain = [_txbody_lst(shape)]
    layout = slide.slide_layout
    master = layout.slide_master
    tx_styles = master._element.find("p:txStyles", NS)

    if shape.is_placeholder:
        pf = shape.placeholder_format
        lph = next((p for p in layout.placeholders if p.placeholder_format.idx == pf.idx), None)
        if lph is None:
            lph = next((p for p in layout.placeholders if p.placeholder_format.type == pf.type), None)
        if lph is not None:
            chain.append(_txbody_lst(lph))

        if pf.type in TITLE_TYPES:
            want, style = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}, "p:titleStyle"
        elif pf.type in FOOTER_TYPES:
            want, style = {pf.type}, "p:otherStyle"
        else:
            want, style = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, pf.type}, "p:bodyStyle"
        mph = next((p for p in master.placeholders if p.placeholder_format.type in want), None)
        if mph is not None:
            chain.append(_txbody_lst(mph))
        if tx_styles is not None:
            chain.append(tx_styles.find(style, NS))
    else:
        if tx_styles is not None:
            chain.append(tx_styles.find("p:otherStyle", NS))

    chain.append(prs._element.find("p:defaultTextStyle", NS))
    return chain


def _font_scale(text_frame):
    body_pr = text_frame._txBody.find("a:bodyPr", NS)
    if body_pr is None:
        return 1.0
    auto = body_pr.find("a:normAutofit", NS)
    if auto is not None and auto.get("fontScale"):
        return int(auto.get("fontScale")) / 100000
    return 1.0


def _run_size(run, para, chain, default=18.0):
    if run.font.size is not None:
        return run.font.size.pt, "explicit"
    if para.font.size is not None:
        return para.font.size.pt, "paragraph"
    for lst in chain:
        sz = _lvl_sz(lst, para.level)
        if sz:
            return sz, "inherited"
    return default, "default"


# -------------------------------------------------------------- traversal --

def _walk(shapes):
    for sh in shapes:
        try:
            st = sh.shape_type
        except NotImplementedError:
            st = None
        if st == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(sh.shapes)
        else:
            yield sh


def _is_smartart(shape):
    gd = shape._element.find(".//a:graphicData", NS)
    return gd is not None and gd.get("uri") == DIAGRAM_URI


def _typefaces(run):
    rpr = run._r.find("a:rPr", NS)
    if rpr is None:
        return []
    out = []
    for tag in ("a:latin", "a:ea"):
        el = rpr.find(tag, NS)
        if el is not None and el.get("typeface") and not el.get("typeface").startswith("+"):
            out.append(el.get("typeface"))
    return out


def analyze(path, min_body=12.0, max_chars=400):
    prs = Presentation(path)
    slides_out = []
    deck_fonts = Counter()
    deck_sizes = Counter()

    for no, slide in enumerate(prs.slides, start=1):
        info = {
            "slide": no,
            "hidden": slide._element.get("show") == "0",
            "layout": slide.slide_layout.name,
            "title": "",
            "text": [],
            "chars": 0,
            "counts": Counter(),
            "tables": [],
            "charts": [],
            "body_sizes": [],
            "all_sizes": [],
            "min_autofit_scale": 1.0,
            "max_bullet_level": 0,
            "has_source_note": False,
            "numeric_claims": [],
            "placeholder_text": [],
        }

        def take_text(tf, chain, is_title=False, is_footer=False):
            scale = _font_scale(tf)
            info["min_autofit_scale"] = min(info["min_autofit_scale"], scale)
            for para in tf.paragraphs:
                ptext = "".join(r.text for r in para.runs).strip()
                if not ptext:
                    continue
                info["text"].append(ptext)
                info["max_bullet_level"] = max(info["max_bullet_level"], para.level)
                is_note = bool(SOURCE_RE.search(ptext))
                if is_note:
                    info["has_source_note"] = True
                info["numeric_claims"] += NUMBER_RE.findall(ptext)
                info["placeholder_text"] += PLACEHOLDER_TEXT_RE.findall(ptext)
                for run in para.runs:
                    n = len(re.sub(r"\s", "", run.text))
                    if not n:
                        continue
                    size, _src = _run_size(run, para, chain)
                    size = round(size * scale, 1)
                    info["all_sizes"].append((size, n))
                    deck_sizes[size] += n
                    if not (is_title or is_footer or is_note):
                        info["body_sizes"].append((size, n))
                    for tf_name in _typefaces(run):
                        deck_fonts[tf_name] += n

        for sh in _walk(slide.shapes):
            ptype = _ph_type(sh) if sh.is_placeholder else None
            is_title = ptype in TITLE_TYPES
            is_footer = ptype in FOOTER_TYPES
            try:
                st = sh.shape_type
            except NotImplementedError:
                st = None

            if getattr(sh, "has_table", False) and sh.has_table:
                tbl = sh.table
                info["counts"]["table"] += 1
                info["tables"].append({"rows": len(tbl.rows), "cols": len(tbl.columns)})
                for row in tbl.rows:
                    for cell in row.cells:
                        take_text(cell.text_frame, [prs._element.find("p:defaultTextStyle", NS)])
                continue
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                info["counts"]["chart"] += 1
                try:
                    ctype = str(ch.chart_type).split(".")[-1].split(" ")[0]
                except Exception:
                    ctype = "unknown"
                try:
                    n_series = len(list(ch.plots[0].series)) if len(ch.plots) else 0
                except Exception:
                    n_series = None
                info["charts"].append({
                    "type": ctype,
                    "has_title": bool(ch.has_title),
                    "series": n_series,
                })
                continue
            if _is_smartart(sh):
                info["counts"]["smartart"] += 1
                continue
            if st in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE) or (ptype == PP_PLACEHOLDER.PICTURE and not sh.has_text_frame):
                info["counts"]["picture"] += 1
                continue
            if st in (MSO_SHAPE_TYPE.LINE,) or sh._element.tag.endswith("}cxnSp"):
                info["counts"]["connector"] += 1
                continue

            if sh.has_text_frame and sh.text_frame.text.strip():
                if is_title and not info["title"]:
                    info["title"] = sh.text_frame.text.strip().replace("\n", " ")
                chain = _inherited_chain(sh, slide, prs)
                take_text(sh.text_frame, chain, is_title, is_footer)
                if st == MSO_SHAPE_TYPE.AUTO_SHAPE:
                    info["counts"]["autoshape"] += 1
                elif not sh.is_placeholder:
                    info["counts"]["textbox"] += 1
            elif st == MSO_SHAPE_TYPE.AUTO_SHAPE:
                info["counts"]["autoshape"] += 1

        if slide.has_notes_slide:
            info["notes"] = slide.notes_slide.notes_text_frame.text.strip()[:500]
        slides_out.append(_finish(info, min_body, max_chars))

    families = [f for f, _ in deck_fonts.most_common()]
    deck_flags = []
    if len(families) > 3:
        deck_flags.append(f"폰트 {len(families)}종 혼용: {', '.join(families[:6])}")
    flagged = [s["slide"] for s in slides_out if s["flags"]]

    return {
        "file": str(path),
        "slide_count": len(slides_out),
        "slide_size_in": [round(prs.slide_width / 914400, 2), round(prs.slide_height / 914400, 2)],
        "font_families": dict(deck_fonts.most_common()),
        "font_size_histogram_pt": {str(k): v for k, v in sorted(deck_sizes.items())},
        "deck_flags": deck_flags,
        "flagged_slides": flagged,
        "slides": slides_out,
    }


def _weighted_median(pairs):
    if not pairs:
        return None
    expanded = []
    for size, n in pairs:
        expanded += [size] * n
    return statistics.median(expanded)


def _finish(info, min_body, max_chars):
    body = info.pop("body_sizes")
    allsz = info.pop("all_sizes")
    info["chars"] = sum(len(re.sub(r"\s", "", t)) for t in info["text"])
    info["min_body_pt"] = min((s for s, _ in body), default=None)
    info["median_body_pt"] = _weighted_median(body)
    info["distinct_sizes"] = sorted({s for s, _ in allsz})
    small = sum(n for s, n in body if s < min_body)
    total = sum(n for _, n in body)
    info["small_text_ratio"] = round(small / total, 2) if total else 0.0
    info["autofit_scale"] = info.pop("min_autofit_scale")
    counts = info.pop("counts")
    info["counts"] = dict(counts)
    visuals = sum(counts[k] for k in ("table", "chart", "picture", "smartart", "autoshape"))
    info["numeric_claims"] = sorted(set(info["numeric_claims"]))[:20]
    info["placeholder_text"] = sorted(set(info["placeholder_text"]))

    flags = []
    if info["min_body_pt"] is not None and info["min_body_pt"] < 10:
        flags.append(f"본문 최소 {info['min_body_pt']}pt (판독 곤란)")
    elif info["min_body_pt"] is not None and info["min_body_pt"] < min_body:
        flags.append(f"본문 최소 {info['min_body_pt']}pt (작음)")
    if info["autofit_scale"] < 0.8:
        flags.append(f"자동 맞춤으로 {int(info['autofit_scale'] * 100)}%까지 축소")
    if info["chars"] > max_chars:
        flags.append(f"텍스트 과밀 ({info['chars']}자)")
    if visuals == 0 and info["chars"] > 150:
        flags.append("시각요소 없이 텍스트만")
    if info["max_bullet_level"] >= 3:
        flags.append(f"글머리 {info['max_bullet_level'] + 1}단계")
    if len(info["distinct_sizes"]) >= 5:
        flags.append(f"글자 크기 {len(info['distinct_sizes'])}종 혼재")
    if info["numeric_claims"] and not info["has_source_note"]:
        flags.append("출처 없는 수치")
    for c in info["charts"]:
        if not c["has_title"]:
            flags.append(f"차트 제목 없음 ({c['type']})")
    for t in info["tables"]:
        if t["cols"] > 7 or t["rows"] > 12:
            flags.append(f"큰 표 {t['rows']}x{t['cols']}")
    if info["placeholder_text"]:
        flags.append("미완성 텍스트: " + ", ".join(info["placeholder_text"]))
    if not info["title"] and not info["hidden"]:
        flags.append("제목 개체 없음 (텍스트 상자로 제목을 썼는지 이미지로 확인)")
    if info["hidden"]:
        flags.append("숨김 슬라이드")
    info["flags"] = flags
    return info


def to_markdown(result):
    lines = [
        f"# PPTX 분석: {result['file']}",
        "",
        f"- 슬라이드 {result['slide_count']}장, 크기 {result['slide_size_in'][0]}x{result['slide_size_in'][1]} in",
        f"- 폰트: {', '.join(f'{k}({v})' for k, v in list(result['font_families'].items())[:6]) or '테마 기본'}",
        f"- 플래그 슬라이드: {', '.join(map(str, result['flagged_slides'])) or '없음'}",
    ]
    for f in result["deck_flags"]:
        lines.append(f"- ⚠ {f}")
    lines += [
        "",
        "| # | 제목 | 글자 수 | 본문 최소/중앙 pt | 표 | 차트 | 그림 | 도형 | 플래그 |",
        "|---|---|---|---|---|---|---|---|---|",
    ]
    for s in result["slides"]:
        c = s["counts"]
        title = (s["title"] or "-")[:40].replace("|", "/")
        lines.append(
            f"| {s['slide']} | {title} | {s['chars']} | {s['min_body_pt']}/{s['median_body_pt']} | "
            f"{c.get('table', 0)} | {c.get('chart', 0)} | {c.get('picture', 0)} | "
            f"{c.get('autoshape', 0) + c.get('smartart', 0)} | {'; '.join(s['flags'])} |"
        )
    return "\n".join(lines) + "\n"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx")
    ap.add_argument("-o", "--output", help="JSON 저장 경로 (생략하면 stdout)")
    ap.add_argument("--md", help="요약 Markdown 저장 경로")
    ap.add_argument("--min-body", type=float, default=12.0, help="본문 최소 글자 크기 기준 pt (RFP 기준이 있으면 그 값)")
    ap.add_argument("--max-chars", type=int, default=400, help="슬라이드당 과밀 판단 글자 수")
    args = ap.parse_args()

    result = analyze(args.pptx, args.min_body, args.max_chars)
    data = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        with open(args.output, "w", encoding="utf-8") as f:
            f.write(data)
    else:
        print(data)
    if args.md:
        with open(args.md, "w", encoding="utf-8") as f:
            f.write(to_markdown(result))
    print(f"분석 완료: {result['slide_count']}장, 플래그 {len(result['flagged_slides'])}장", file=sys.stderr)


if __name__ == "__main__":
    main()
